"""產生「網站內容」試算表（xlsx，上傳後轉成 Google 試算表）與備份資料 data/content.json。
工作表名稱與欄位名稱必須和 assets/app.js 讀取的一致。
用法：python tools/build_sheet.py [輸出的 xlsx 路徑]
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
d = date.fromisoformat
DATE_COLS = ("日期", "第一天", "第二天")

SHEETS = {
    "設定": (["項目", "內容", "說明"], [
        ["標語", "數據治理．決策有理", "開場標題下方的一句話"],
        ["平台網址", "", "AI 分析平台的網址；空白時按鈕會捲到平台介紹"],
        ["聯絡窗口", "", "例如：教育資料素養提升計畫辦公室"],
        ["聯絡信箱", "", ""],
        ["聯絡電話", "", ""],
    ]),
    "消息": (["日期", "分類", "標題", "內容", "連結", "顯示"], [
        [d("2026-09-30"), "研習", "115 年學校行政講師培訓四區場次公布", "10/31 中區首場開始，請各縣市完成學員推薦。", "", "是"],
        [d("2026-09-28"), "計畫", "初階與學校行政教材將進行專家審查", "10/5 起兩輪委員書審，10/23 召開共識會議定稿。", "", "是"],
        [d("2026-09-10"), "平台", "AI 輔助分析平台改版", "新增資料去識別化，並可用點選或對話方式操作。", "", "是"],
    ]),
    "簡章": (["代號", "顯示", "分頁名稱", "簡章標題", "年度", "狀態文字", "簡章PDF"], [
        ["T115-ADM", "是", "學校行政講師培訓", "教育資料素養講師培訓工作坊（學校行政）簡章", 115, "縣市推薦中", ""],
        ["T116-TCH", "否", "學校教師講師培訓", "教育資料素養講師培訓工作坊（學校教師）簡章", 116, "縣市推薦中", ""],
        ["C116-CITY", "否", "縣市行政研習", "教育資料素養增能暨應用研習（縣市行政）簡章", 116, "即將開放", ""],
    ]),
    "簡章條目": (["簡章代號", "條目名稱", "內容"], [
        ["T115-ADM", "參加對象", "由各縣市推薦之學校行政講師人選，可為跨局處行政人員或校長。"],
        ["T115-ADM", "參加資格", "已完成 A1、A2 及 B5-1 研習，報名時請檢附證明。"],
        ["T115-ADM", "報名方式", "採縣市推薦制，請洽所屬縣市數位學習推動辦公室，由縣市統一報名。"],
        ["T115-ADM", "課程安排", "・第一天：初階「資料素養基礎與解讀」3 小時＋講師培訓前置 3 小時\n・第二天：進階「教育數據分析與應用」6 小時"],
        ["T115-ADM", "場次與地點", "【場次表】"],
        ["T115-ADM", "講師資格認證", "1. 完成兩天培訓（共 12 小時）\n2. 繳交公版簡報及試教影片，由計畫團隊審查\n3. 審查通過取得講師資格，效期 2 年\n4. 參與講師實作交流會"],
        ["T116-TCH", "參加對象", "由各縣市或學校推薦之數位教學優良教師。"],
        ["T116-TCH", "參加資格", "已完成 A1、A2 及 B5-1 研習，報名時請檢附證明。"],
        ["T116-TCH", "報名方式", "採縣市或學校推薦制，請洽所屬縣市數位學習推動辦公室。"],
        ["T116-TCH", "課程安排", "・第一天：初階「資料素養基礎與解讀」3 小時＋講師培訓前置 3 小時\n・第二天：進階「教育數據分析與應用」6 小時（學校教師）"],
        ["T116-TCH", "場次與地點", "【場次表】"],
        ["T116-TCH", "講師資格認證", "1. 完成兩天培訓（共 12 小時）\n2. 繳交公版簡報及試教影片，由計畫團隊審查\n3. 審查通過取得講師資格，效期 2 年\n4. 參與講師實作交流會"],
        ["C116-CITY", "參加對象", "縣市教育局處及數位學習推動辦公室相關人員。"],
        ["C116-CITY", "參加資格", "已完成 A1、A2 及 B5-1 研習。"],
        ["C116-CITY", "課程安排", "・第一天：初階「資料素養基礎與解讀」3 小時\n・第二天：進階「教育數據分析與應用」6 小時，實際分析縣市教育數據\n・參與縣市成果分享會"],
        ["C116-CITY", "場次與地點", "【場次表】"],
    ]),
    "場次": (["簡章代號", "區域", "第一天", "第二天", "時間", "地點", "報名表連結", "顯示"], [
        ["T115-ADM", "中區", d("2026-10-31"), d("2026-11-21"), "", "", "", "是"],
        ["T115-ADM", "北區", d("2026-11-01"), d("2026-11-22"), "", "", "", "是"],
        ["T115-ADM", "南區", d("2026-11-07"), d("2026-12-05"), "", "", "", "是"],
        ["T115-ADM", "東區", d("2026-11-08"), d("2026-11-29"), "", "", "", "是"],
    ]),
    "活動": (["日期", "名稱", "說明", "顯示"], [
        [d("2026-12-09"), "縣市分享會", "配合自主學習節，至少 4 個縣市分享數據分析的做法與成果。", "是"],
        [d("2026-12-23"), "跨區共通講師交流會", "下午場，世界咖啡館形式，討論不同資料來源、不同樣本數的分析方式，以及地方特色與數據分析的關聯。", "是"],
        [d("2026-12-29"), "輔導團共識會議（暫定）", "配合 22 縣市工作會議，分享研習內容並交流數據分析情形。", "是"],
    ]),
    "常見問題": (["問題", "回答", "顯示"], [
        ["我可以自己報名嗎？", "115 年的講師培訓採縣市推薦制，請洽所屬縣市數位學習推動辦公室。116 年起由縣市辦理的學校行政與教師研習改採申請制，屆時會在本頁公告。", "是"],
        ["參加前需要先修哪些研習？", "需先完成 A1、A2 及 B5-1 研習，內容包含 AI 基礎操作與數位學習理論。", "是"],
        ["怎麼取得講師資格？", "1. 完成兩天培訓\n2. 繳交公版簡報的教學演示影片，由計畫團隊審查\n3. 通過後取得講師資格，效期 2 年\n4. 期滿參加回流研習，並出示期間內擔任講師的證明即可延續", "是"],
        ["116 年會有哪些研習？", "預計辦理縣市行政研習 10 場、學校行政講師培訓 4 場、學校教師講師培訓 6 場，場次確定後會在最新消息公告。", "是"],
    ]),
}

GUIDE = [
    ("這份試算表是「教育資料素養提升計畫」網站的內容後台。", True),
    ("改完內容後，重新整理網頁就會看到新內容（瀏覽器有快取時，最多等 1 分鐘）。", False),
    ("", False),
    ("【基本規則】", True),
    ("1. 每張工作表的第一列是欄位名稱，請不要修改、刪除或移動第一列，也不要改工作表名稱。", False),
    ("2. 一列就是一筆資料。新增一筆就往下加一列；中間不要留空白列。", False),
    ("3. 「顯示」欄選「是」才會出現在網頁上。過期的資料改成「否」即可，不需要刪除。", False),
    ("4. 日期請填 2026-10-31 這種格式（年-月-日）。", False),
    ("5. 連結請貼完整網址（https:// 開頭）。雲端硬碟的 PDF 請先設為「知道連結的使用者可檢視」。", False),
    ("6. 網頁要讀得到這份試算表，它必須設為「知道連結的使用者可檢視」，所以請勿填寫個資或內部資料。", False),
    ("7. 改錯了：上方選單「檔案 > 版本記錄」可以還原到先前版本。", False),
    ("", False),
    ("【各工作表】", True),
    ("設定：標語、平台網址、聯絡資訊。只改「內容」欄。聯絡資訊空白時，網頁不會顯示該列。", False),
    ("消息：網頁顯示最新 5 則（依日期排序）。", False),
    ("簡章：一列一份簡章。「代號」是這份簡章的識別碼，簡章條目與場次都用它對應，建立後請不要改。同時有兩份以上顯示為「是」時，網頁會出現分頁切換。", False),
    ("簡章條目：依列的先後順序顯示，編號（一、二、三…）由網頁自動產生。", False),
    ("　・內容中每行以「・」開頭會變成條列；以「1.」「2.」開頭會變成編號清單。", False),
    ("　・內容填「【場次表】」，會在該條目的位置插入場次表；沒有填的話，場次表會放在最後。", False),
    ("場次：「區域」含北、中、南、東時會自動套用 LOGO 對應顏色。狀態（幾天後開始、進行中、已結束）由網頁依日期自動計算，平常顯示簡章的「狀態文字」。有填報名表連結時會顯示報名連結。", False),
    ("活動：交流會、分享會等單場活動，依日期排序。", False),
    ("常見問題：依列的先後順序顯示。", False),
    ("", False),
    ("【範本】", True),
    ("簡章中的「T116-TCH 學校教師講師培訓」與「C116-CITY 縣市行政研習」是依 0910 簡報擬的範本，內容待確認，目前顯示為「否」。", False),
]

WIDTHS = {"日期": 13, "第一天": 13, "第二天": 13, "顯示": 7, "代號": 12, "簡章代號": 12, "年度": 7,
          "區域": 8, "分類": 8, "內容": 60, "說明": 44, "回答": 70, "標題": 36, "簡章標題": 44,
          "問題": 30, "名稱": 26, "條目名稱": 16, "地點": 30, "時間": 14}
LAST_ROW = 500


def build_xlsx(path):
    wb = Workbook()
    guide = wb.active
    guide.title = "使用說明"
    guide.column_dimensions["A"].width = 120
    for i, (text, bold) in enumerate(GUIDE, start=1):
        c = guide.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=12 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    head_fill = PatternFill("solid", fgColor="EAF5F7")
    for name, (headers, rows) in SHEETS.items():
        s = wb.create_sheet(name)
        s.append(headers)
        for r in rows:
            s.append(r)
        s.freeze_panes = "A2"
        for j, h in enumerate(headers, start=1):
            head = s.cell(row=1, column=j)
            head.font = Font(bold=True)
            head.fill = head_fill
            col = head.column_letter
            s.column_dimensions[col].width = WIDTHS.get(h, 20)
            for i in range(2, len(rows) + 2):
                s.cell(row=i, column=j).alignment = Alignment(wrap_text=True, vertical="top")
            if h in DATE_COLS:
                for i in range(2, LAST_ROW + 1):
                    s.cell(row=i, column=j).number_format = "yyyy-mm-dd"
                dv = DataValidation(type="date", operator="between",
                                    formula1="DATE(2025,1,1)", formula2="DATE(2030,12,31)",
                                    allow_blank=True, showErrorMessage=True,
                                    errorTitle="日期格式", error="請填日期，例如 2026-10-31")
                dv.add(f"{col}2:{col}{LAST_ROW}")
                s.add_data_validation(dv)
            if h == "顯示":
                dv = DataValidation(type="list", formula1='"是,否"', allow_blank=False,
                                    showErrorMessage=True, errorTitle="請選擇", error="請選「是」或「否」")
                dv.add(f"{col}2:{col}{LAST_ROW}")
                s.add_data_validation(dv)
    wb.save(path)


def cell_text(v):
    if isinstance(v, date):
        return v.isoformat()
    return "" if v is None else str(v)


def build_json(path):
    out = {name: [dict(zip(headers, map(cell_text, r))) for r in rows]
           for name, (headers, rows) in SHEETS.items()}
    path.parent.mkdir(exist_ok=True)
    path.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "網站內容.xlsx"
    build_xlsx(xlsx)
    build_json(ROOT / "data" / "content.json")
    print("ok", xlsx)
